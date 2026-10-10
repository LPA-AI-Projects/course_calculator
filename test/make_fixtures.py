# Generates differential-test fixtures: random inputs + random rate settings are written into copies
# of the original Excel, recalculated by LibreOffice, and the Excel results saved next to the config used.
import openpyxl, random, json, subprocess, os, sys, shutil, warnings
warnings.filterwarnings("ignore")
SRC, WORK = sys.argv[1], sys.argv[2]
random.seed(42)
os.makedirs(WORK, exist_ok=True)
base = openpyxl.load_workbook(SRC)
courses = [r[0] for r in base['Courses'].iter_rows(min_row=2, max_col=1, values_only=True)]
N = 90
fx = []
for i in range(N):
    wb = openpyxl.load_workbook(SRC)
    R, C = wb['Rates'], wb['Calculator']
    cfg = {"settings": {}, "categories": [], "modes": [], "levels": [], "hoursPerDay": []}
    for r in range(5, 18):
        lo_o = random.randint(80, 200); hi_o = lo_o + random.randint(0, 80)
        lo_f = lo_o + random.randint(0, 60); hi_f = lo_f + random.randint(0, 80)
        for col, v in zip((2, 3, 5, 6), (lo_o, hi_o, lo_f, hi_f)): R.cell(r, col).value = v
        cfg["categories"].append({"name": R.cell(r, 1).value, "onlineLow": lo_o, "onlineHigh": hi_o, "offlineLow": lo_f, "offlineHigh": hi_f})
    mult = random.choice([2.8, 3.0, 3.15, 2.5]); vat = random.choice([0.05, 0.1, 0.0]); rnd = random.choice([1, 5, 10, 0.5, 25])
    R['B21'], R['B22'], R['B23'] = mult, vat, rnd
    cfg["settings"] = {"multiplier": mult, "vatRate": vat, "roundTo": rnd}
    for r in range(27, 30):
        m = random.choice([1, 1, 1.1, 1.2, 0.9]); R.cell(r, 3).value = m
        cfg["modes"].append({"name": R.cell(r, 1).value, "rateCard": R.cell(r, 2).value, "multiplier": m})
    for r in range(32, 35):
        m = random.choice([1, 0.9, 1.15, 1.3]); R.cell(r, 2).value = m
        cfg["levels"].append({"name": R.cell(r, 1).value, "multiplier": m})
    for r in range(46, 49):
        m = random.choice([1, 1.25, 1.1, 0.95]); R.cell(r, 2).value = m
        cfg["hoursPerDay"].append({"hours": R.cell(r, 1).value, "multiplier": m})
    inp = {"course": random.choice(courses), "level": random.choice(["Beginner", "Intermediate", "Advanced"]),
           "mode": random.choice(["Online", "Classroom", "Onsite"]),
           "days": random.choice([1, 2, 3, 4, 5, 7, 10, 14, 20, 30, 45, 90]), "hoursPerDay": random.choice([2, 4, 8]),
           "overrideHours": random.choice([None, None, None, 6, 12.5, 40])}
    C['C4'], C['C5'], C['C6'], C['C7'], C['C8'], C['C9'] = inp["course"], inp["level"], inp["mode"], inp["days"], inp["hoursPerDay"], inp["overrideHours"]
    wb.save(f"{WORK}/f{i}.xlsx")
    fx.append({"config": cfg, "input": inp})
subprocess.run(f"cd {WORK} && soffice --headless --calc --convert-to xlsx:'Calc MS Excel 2007 XML' --outdir out f*.xlsx >/dev/null 2>&1", shell=True)
cat_of = {r[0]: r[1] for r in base['Courses'].iter_rows(min_row=2, max_col=2, values_only=True)}
for i, f in enumerate(fx):
    w = openpyxl.load_workbook(f"{WORK}/out/f{i}.xlsx", data_only=True)['Calculator']
    f["expected"] = {"category": w['C12'].value, "rate": w['C13'].value, "hours": w['C14'].value, "trainerFee": w['C15'].value,
                     "courseFee": w['C17'].value, "vat": w['C18'].value, "total": w['C20'].value}
    f["categoryOf"] = cat_of[f["input"]["course"]]
json.dump(fx, open("test/fixtures.json", "w"), ensure_ascii=False)
print("fixtures:", len(fx), "sample:", fx[0]["expected"])
